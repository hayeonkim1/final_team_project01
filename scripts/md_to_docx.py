"""Convert markdown file to Word document."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement


def set_cell_shading(cell, color_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_formatted_runs(paragraph, text: str):
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        chunk = match.group(0)
        if chunk.startswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("*"):
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def is_table_row(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and "|" in stripped[1:-1]


def is_separator_row(line: str) -> bool:
    return bool(re.match(r"^\|\s*[-:]+(\s*\|\s*[-:]+)+\s*\|?\s*$", line.strip()))


def parse_table_row(line: str) -> list[str]:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def convert(md_path: Path, docx_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(10.5)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=3)
            i += 1
            continue

        if is_table_row(stripped):
            table_lines = []
            while i < len(lines) and is_table_row(lines[i].strip()):
                table_lines.append(lines[i].strip())
                i += 1

            data_rows = [r for r in table_lines if not is_separator_row(r)]
            if not data_rows:
                continue

            col_count = len(parse_table_row(data_rows[0]))
            table = doc.add_table(rows=len(data_rows), cols=col_count)
            table.style = "Table Grid"

            for r_idx, row_line in enumerate(data_rows):
                cells = parse_table_row(row_line)
                for c_idx in range(col_count):
                    cell_text = cells[c_idx] if c_idx < len(cells) else ""
                    cell = table.rows[r_idx].cells[c_idx]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    add_formatted_runs(p, cell_text)
                    if r_idx == 0:
                        set_cell_shading(cell, "E8E8E8")
                        for run in p.runs:
                            run.bold = True
            doc.add_paragraph()
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            run = p.add_run(stripped[2:].strip())
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            add_formatted_runs(p, re.sub(r"^\d+\.\s", "", stripped))
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, stripped[2:].strip())
            i += 1
            continue

        if stripped.startswith("   - "):
            p = doc.add_paragraph(style="List Bullet 2")
            add_formatted_runs(p, stripped[5:].strip())
            i += 1
            continue

        p = doc.add_paragraph()
        add_formatted_runs(p, stripped)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"Created: {docx_path}")


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    md = root / "docs" / "짐토리_프로젝트_종합정리.md"
    out = root / "docs" / "짐토리_프로젝트_종합정리.docx"
    if len(sys.argv) > 1:
        md = Path(sys.argv[1])
    if len(sys.argv) > 2:
        out = Path(sys.argv[2])
    convert(md, out)
